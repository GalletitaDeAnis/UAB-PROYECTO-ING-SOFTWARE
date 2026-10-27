"""Genera los documentos de Word a partir de los Markdown de entregables/.

El Markdown del repositorio es la fuente; los .docx se regeneran con:  python generar_word.py
Así el Word que se entrega nunca queda desactualizado respecto del código.
"""

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AQUI = Path(__file__).resolve().parent
ENTREGABLES = AQUI.parent

INTEGRANTES = [
    "Paul Steffano Velarde Chuquimia",
    "Paolo Pizarro A.",
    "Xander Boris Pinto Montes",
    "Luis Alberto Mamani L.",
]
AZUL = RGBColor(0x1A, 0x52, 0x76)


def preparar(doc: Document) -> None:
    estilo = doc.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(10.5)
    for nombre, tam in (("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12)):
        doc.styles[nombre].font.color.rgb = AZUL
        doc.styles[nombre].font.size = Pt(tam)
    for seccion in doc.sections:
        seccion.left_margin = seccion.right_margin = Cm(2)
        seccion.top_margin = seccion.bottom_margin = Cm(2)


def caratula(doc: Document, titulo: str) -> None:
    def linea(texto, tam=12, negrita=False, espacio=6):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(espacio)
        r = p.add_run(texto)
        r.font.size = Pt(tam)
        r.bold = negrita
        return p

    linea("", espacio=40)
    linea("UNIVERSIDAD ADVENTISTA DE BOLIVIA", 18, True)
    linea("UAB", 16, True, 30)
    linea("Facultad / Carrera: Ingeniería de Sistemas", 12)
    linea("Materia: Ingeniería de Software", 12, espacio=30)
    linea("PROYECTO «UNIÓN BOLIVIANA»", 16, True)
    linea("Sistema de Reservas y Pensionamiento de Almuerzos", 13)
    linea(titulo, 14, True, 40)
    linea("Integrantes del grupo", 12, True)
    for nombre in INTEGRANTES:
        linea(nombre, 12, espacio=2)
    linea("", espacio=20)
    linea("Docente: Ing. Roy Carrasco", 12)
    linea("Aplicación: https://almuerzos-union-boliviana.onrender.com", 10)
    linea("Repositorio: https://github.com/GalletitaDeAnis/UAB-PROYECTO-ING-SOFTWARE (rama desarrollo)", 10)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|<https?://[^>]+>)")


def texto_en_linea(parrafo, texto: str) -> None:
    for trozo in INLINE.split(texto):
        if not trozo:
            continue
        if trozo.startswith("**") and trozo.endswith("**"):
            parrafo.add_run(trozo[2:-2]).bold = True
        elif trozo.startswith("`") and trozo.endswith("`"):
            r = parrafo.add_run(trozo[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
        elif trozo.startswith("["):
            r = parrafo.add_run(re.match(r"\[([^\]]+)\]", trozo).group(1))
            r.font.color.rgb = AZUL
        elif trozo.startswith("<http"):
            r = parrafo.add_run(trozo[1:-1])
            r.font.color.rgb = AZUL
        else:
            parrafo.add_run(trozo)


def sombrear(celda, color: str) -> None:
    tc = celda._tc.get_or_add_tcPr()
    sombra = OxmlElement("w:shd")
    sombra.set(qn("w:val"), "clear")
    sombra.set(qn("w:color"), "auto")
    sombra.set(qn("w:fill"), color)
    tc.append(sombra)


def tabla(doc: Document, filas: list[list[str]]) -> None:
    columnas = max(len(f) for f in filas)
    t = doc.add_table(rows=0, cols=columnas)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, fila in enumerate(filas):
        celdas = t.add_row().cells
        for j in range(columnas):
            celda = celdas[j]
            celda.paragraphs[0].paragraph_format.space_after = Pt(0)
            texto_en_linea(celda.paragraphs[0], fila[j] if j < len(fila) else "")
            for r in celda.paragraphs[0].runs:
                r.font.size = Pt(9)
                if i == 0:
                    r.bold = True
            if i == 0:
                sombrear(celda, "D6EAF8")
    doc.add_paragraph()


def celdas(linea: str) -> list[str]:
    return [c.strip() for c in linea.strip().strip("|").split("|")]


def markdown_a_word(doc: Document, ruta: Path, nivel_base: int = 0) -> None:
    lineas = ruta.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lineas):
        linea = lineas[i]
        if linea.startswith("```"):
            i += 1
            codigo = []
            while i < len(lineas) and not lineas[i].startswith("```"):
                codigo.append(lineas[i])
                i += 1
            p = doc.add_paragraph()
            r = p.add_run("\n".join(codigo))
            r.font.name = "Consolas"
            r.font.size = Pt(8.5)
            i += 1
            continue
        if linea.startswith("|"):
            filas = []
            while i < len(lineas) and lineas[i].startswith("|"):
                if not re.match(r"^\|\s*:?-{3,}", lineas[i]):
                    filas.append(celdas(lineas[i]))
                i += 1
            tabla(doc, filas)
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", linea)
        if m:
            nivel = min(len(m.group(1)) + nivel_base, 4)
            doc.add_heading(re.sub(r"[*`]", "", m.group(2)), level=nivel)
        elif imagen := re.match(r"!\[([^\]]*)\]\(([^)]+)\)", linea.strip()):
            archivo = (ruta.parent / imagen.group(2)).resolve()
            if archivo.exists():
                doc.add_picture(str(archivo), width=Cm(16.5))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                pie = doc.add_paragraph(imagen.group(1))
                pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pie.runs[0].italic = True
        elif re.match(r"^\s*[-*]\s+", linea):
            texto = re.sub(r"^\s*[-*]\s+(\[[ x]\]\s+)?", "", linea)
            marca = "☐ " if "[ ]" in linea[:8] else ""
            p = doc.add_paragraph(style="List Bullet")
            texto_en_linea(p, marca + texto)
        elif re.match(r"^\s*\d+\.\s+", linea):
            p = doc.add_paragraph(style="List Number")
            texto_en_linea(p, re.sub(r"^\s*\d+\.\s+", "", linea))
        elif linea.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            texto_en_linea(p, linea.lstrip("> "))
            for r in p.runs:
                r.italic = True
        elif linea.strip():
            texto_en_linea(doc.add_paragraph(), linea.strip())
        i += 1


def documento(nombre: str, titulo: str, fuentes: list[str]) -> Path:
    doc = Document()
    preparar(doc)
    caratula(doc, titulo)
    for indice, fuente in enumerate(fuentes):
        if indice:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        markdown_a_word(doc, ENTREGABLES / fuente)
    salida = AQUI / nombre
    doc.save(salida)
    return salida


if __name__ == "__main__":
    generados = [
        documento("Caratula_Union_Boliviana_UAB.docx", "Presentación del proyecto académico", []),
        documento("Backlog_e_historias_de_usuario.docx", "Backlog e historias de usuario (versión final)", ["backlog.md"]),
        documento("Casos_de_uso_y_diagrama_de_clases.docx", "Casos de uso y diagrama de clases", ["casos-de-uso.md", "diagrama-de-clases.md"]),
        documento("Arquitectura_y_base_de_datos.docx", "Arquitectura y decisiones de diseño", ["arquitectura.md", "base-de-datos.md", "api-graphql.md"]),
        documento("Testing_DoD_y_refactorizacion.docx", "Testing, Definition of Done y refactorización", ["definition-of-done.md", "registro-de-refactorizacion.md"]),
        documento("Entregas_retros_y_uso_de_IA.docx", "Entregas, retrospectivas y uso de IA", ["entregas.md", "acta-planning-poker.md", "plantilla-retro.md", "declaracion-uso-de-ia.md"]),
        documento(
            "Documento_completo_del_proyecto.docx",
            "Documento completo del proyecto",
            [
                "backlog.md", "casos-de-uso.md", "diagrama-de-clases.md", "arquitectura.md", "base-de-datos.md",
                "api-graphql.md", "definition-of-done.md", "registro-de-refactorizacion.md", "entregas.md",
                "declaracion-uso-de-ia.md",
            ],
        ),
    ]
    for archivo in generados:
        print(archivo.name)
